"""Fill the Pan / Country / FULL columns of a one2fan.com_g2 settlement report.

Pan     - ISO2 country of the card issuer, determined from the first 6 digits
          (BIN) of the masked card number that the column originally holds.
          BIN -> country reference data: bin-list-data.csv from
          https://github.com/venelinkochev/bin-list-data (374,788 BINs).
          If a BIN is absent from that dataset, the report's own Country value
          is used as a fallback.
Country - left as delivered by the processor.
FULL    - full country name of the Country code.
"""
import csv, json, openpyxl
from openpyxl.styles import Font

SRC = 'work.xlsx'      # copy of the original processor report
OUT = 'out.xlsx'       # filled workbook

# Full country names. Spellings follow the client's existing report convention
# (see reference sheet: "Marocco", "Chili"), standard English elsewhere.
NAMES = {
    'AE': 'United Arab Emirates', 'AR': 'Argentina', 'AT': 'Austria', 'AU': 'Australia',
    'BD': 'Bangladesh', 'BE': 'Belgium', 'BG': 'Bulgaria', 'BR': 'Brazil',
    'CA': 'Canada', 'CH': 'Switzerland', 'CL': 'Chili', 'CN': 'China', 'CO': 'Colombia',
    'CY': 'Cyprus', 'CZ': 'Czech Republic', 'DE': 'Germany', 'DK': 'Denmark',
    'EC': 'Ecuador', 'EE': 'Estonia', 'EG': 'Egypt', 'ES': 'Spain', 'FI': 'Finland',
    'FR': 'France', 'GB': 'United Kingdom', 'GE': 'Georgia', 'GI': 'Gibraltar',
    'GR': 'Greece', 'HK': 'Hong Kong', 'HR': 'Croatia', 'HU': 'Hungary',
    'ID': 'Indonesia', 'IE': 'Ireland', 'IL': 'Israel', 'IN': 'India', 'IT': 'Italy',
    'JP': 'Japan', 'KH': 'Cambodia', 'KR': 'South Korea', 'KW': 'Kuwait',
    'KY': 'Cayman Islands', 'KZ': 'Kazakhstan', 'LT': 'Lithuania', 'LU': 'Luxembourg',
    'LV': 'Latvia', 'MA': 'Marocco', 'MT': 'Malta', 'MX': 'Mexico', 'MY': 'Malaysia',
    'NG': 'Nigeria', 'NL': 'Netherlands', 'NO': 'Norway', 'NZ': 'New Zealand',
    'PA': 'Panama', 'PE': 'Peru', 'PH': 'Philippines', 'PK': 'Pakistan', 'PL': 'Poland',
    'PT': 'Portugal', 'QA': 'Qatar', 'RO': 'Romania', 'RS': 'Serbia', 'RU': 'Russia',
    'SA': 'Saudi Arabia', 'SE': 'Sweden', 'SG': 'Singapore', 'SI': 'Slovenia',
    'SK': 'Slovakia', 'TH': 'Thailand', 'TR': 'Turkey', 'TW': 'Taiwan', 'UA': 'Ukraine',
    'US': 'United States', 'UY': 'Uruguay', 'VN': 'Vietnam', 'ZA': 'South Africa',
}

wb = openpyxl.load_workbook(SRC)
ws = wb['Sheet1']

HDR = 4
last = max(r for r in range(HDR + 1, ws.max_row + 1) if ws.cell(r, 1).value and ws.cell(r, 4).value == 'Payment')

rows = []
for r in range(HDR + 1, last + 1):
    pan = ws.cell(r, 11).value
    rows.append((r, str(pan) if pan else '', ws.cell(r, 12).value))

bins = {p[:6] for _, p, _ in rows if p}
db = {}
with open('bin-list-data.csv') as f:
    for rec in csv.DictReader(f):
        b = rec['BIN']
        if b in bins:
            db[b] = rec['isoCode2'].strip().upper()

report = {}   # what actually got written
mismatch, unresolved, missing_name = [], [], set()

for r, pan, country in rows:
    b = pan[:6]
    iso = db.get(b, '')
    src = 'bin'
    if not iso:                       # BIN not in the reference database
        iso = (country or '').strip().upper()
        src = 'fallback-country'
        if iso:
            unresolved.append((r, b, iso))
    if not iso:
        continue
    if country and iso != country.strip().upper():
        mismatch.append((r, b, iso, country))
    ccode = (country or '').strip().upper() or iso   # FULL expands the Country column
    name = NAMES.get(ccode)
    if name is None:
        missing_name.add(ccode)
        name = ccode
    ws.cell(r, 11).value = iso        # Pan  -> issuing country of the card, from its BIN
    ws.cell(r, 13).value = name       # FULL -> full name of the Country code
    for col in (11, 13):
        ws.cell(r, col).font = Font(name='Arial', size=10)
    report[r] = (iso, country, name, src)

for col, w in (('K', 10), ('L', 10), ('M', 22)):
    ws.column_dimensions[col].width = w

wb.save(OUT)

print(json.dumps({
    'data_rows': len(rows),
    'last_data_row': last,
    'distinct_bins': len(bins),
    'bins_resolved_from_db': len(db),
    'rows_written': len(report),
    'pan_vs_country_mismatch': len(mismatch),
    'bins_without_db_country': sorted({b for _, b, _ in unresolved}),
    'codes_without_full_name': sorted(missing_name),
}, indent=2))
print('\nmismatches (row, bin, pan-country, report-country):')
for m in mismatch:
    print(' ', m)
