"""Fill MOSS workbook with July 2026 data from the CULT workbook.

Source: CULT ' расчет' sheet (= "EXACTLY / Июль", ECB rate 31.07.2026).
Target: MOSS 'exactly 06.2026' sheet -> renamed 'exactly 07.2026'.

Column mapping (CULT ' расчет' -> MOSS):
    D (code)          -> B (code)      : join key
    M (USD Income)    -> D (Income)
    O (USD Chargebacks) -> E (Chargebacks)
    everything else in MOSS is derived by formula from D/E/C and the USD rate.
"""
import shutil
import openpyxl

CULT = '/root/.claude/uploads/031b7d60-2e1e-5f70-9e1e-2280d421c0c0/2720cef6-Cult_2026.07_Exactly.xlsx'
SRC = '/root/.claude/uploads/031b7d60-2e1e-5f70-9e1e-2280d421c0c0/2d4d58aa-MOSS_exactly_07.26.xlsx'
OUT = '/home/user/MyFiles/MOSS_exactly_07.26.xlsx'

# ---- read July source values ------------------------------------------------
cwb = openpyxl.load_workbook(CULT, data_only=True)
cws = cwb[' расчет']

usd_rate = cws['B3'].value            # 1.1485, ECB 31.07.2026
assert cws['B2'].value.strftime('%Y-%m-%d') == '2026-07-31', cws['B2'].value

july = {}                             # ISO code -> (income USD, chargebacks USD)
for r in range(7, 259):
    code = cws.cell(r, 4).value
    if not code:
        continue
    income = cws.cell(r, 13).value or 0      # M
    chargeback = cws.cell(r, 15).value or 0  # O
    july[code] = (income, chargeback)

# ---- write into MOSS --------------------------------------------------------
shutil.copyfile(SRC, OUT)
wb = openpyxl.load_workbook(OUT)
ws = wb['exactly 06.2026']

ws.title = 'exactly 07.2026'
ws['A2'] = 'Exchange rate ECB 31.07.2026'
ws['B3'] = usd_rate
ws['I5'] = 'MOSS 07/26'

EU_FIRST, EU_LAST = 7, 32             # Austria .. Sweden
CZ_ROW = 33
NON_EU_FIRST, NON_EU_LAST = 34, 255

written = 0
for r in range(EU_FIRST, NON_EU_LAST + 1):
    code = ws.cell(r, 2).value
    if not code:
        continue
    income, chargeback = july[code]
    ws.cell(r, 4).value = income                      # D Income (USD)
    written += 1

    if r > CZ_ROW:                                    # non-EU: income column only
        continue

    if chargeback:
        ws.cell(r, 5).value = chargeback              # E Chargebacks
    ws.cell(r, 6).value = f'=D{r}-E{r}'               # F Total Income
    ws.cell(r, 7).value = f'=F{r}/(100+C{r})*100'     # G Net Income
    ws.cell(r, 8).value = f'=G{r}*C{r}/100'           # H VAT
    ws.cell(r, 9).value = f'=G{r}/$B$3'               # I Net Income Total (EUR)
    ws.cell(r, 10).value = f'=H{r}/$B$3'              # J VAT Total (EUR)
    ws.cell(r, 11).value = f'=I{r}+J{r}'              # K Total with Tax

# Total EU row summed from row 8 in F:H, skipping Austria -- start at row 7 like D/I/J/K.
for col in 'EFGH':
    ws[f'{col}257'] = f'=SUM({col}7:{col}32)'

wb.save(OUT)
print(f'rows written: {written}, usd rate: {usd_rate}')
